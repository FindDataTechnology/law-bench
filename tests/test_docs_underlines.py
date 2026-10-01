"""Tests for the underline + slot-style rendering (``add-contract-underlines``)."""

from docx import Document
from docx.oxml.ns import qn

from src.generator import create_docx, find_slots


def _doc(path):
    return Document(str(path))


def _underlined_runs(para):
    """Runs in ``para`` that carry a ``<w:u>`` underline element."""
    out = []
    for r in para.runs:
        rPr = r._r.find(qn("w:rPr"))
        if rPr is not None and rPr.find(qn("w:u")) is not None:
            out.append(r)
    return out


def test_blank_run_real_underline(tmp_path):
    # 5.1: a literal ____ blank becomes an underlined space run, no `_` chars.
    p = create_docx("甲方：____________________", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    assert "_" not in para.text  # no literal underscore characters remain
    assert _underlined_runs(para)  # at least one underlined run


def test_slot_name_underscore_preserved(tmp_path):
    # 5.2: underscores inside a {{slot}} name are not eaten by _split_underlines.
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    assert find_slots(p) == ["party_a"]
    assert "{{party_a}}" in para.text


def test_inline_underline_span(tmp_path):
    # 5.3: _text_ renders the inner text as an underlined run; surroundings not.
    p = create_docx("违约金 _万分之五_ 约定", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    underlined = [r.text for r in _underlined_runs(para)]
    assert "万分之五" in underlined
    for r in para.runs:
        if r.text.strip() and r.text.strip() != "万分之五":
            rPr = r._r.find(qn("w:rPr"))
            assert rPr is None or rPr.find(qn("w:u")) is None


def test_slot_style_underline(tmp_path):
    # 5.4: slot_style="underline" emits a blank, no {{}} text, unfindable.
    p = create_docx("甲方：{{甲方名称}}", tmp_path / "c.docx", slot_style="underline")
    para = _doc(p).paragraphs[0]
    assert "{{" not in para.text and "}}" not in para.text
    assert _underlined_runs(para)
    assert find_slots(p) == []


def test_slot_style_highlight_default(tmp_path):
    # 5.5: default slot_style highlights the token.
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    highlighted = [
        r.text for r in para.runs
        if r._r.find(qn("w:rPr")) is not None
        and r._r.find(qn("w:rPr")).find(qn("w:highlight")) is not None
    ]
    assert "{{party_a}}" in highlighted


def test_slot_style_literal(tmp_path):
    # 5.6: slot_style="literal" keeps token text, no highlight, no underline.
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx", slot_style="literal")
    para = _doc(p).paragraphs[0]
    assert "{{party_a}}" in para.text
    for r in para.runs:
        rPr = r._r.find(qn("w:rPr"))
        assert rPr is None or rPr.find(qn("w:highlight")) is None


def test_heading_spacing(tmp_path):
    # 5.7: Heading 1/2/3 styles have non-zero space_before and space_after.
    p = create_docx("# H1\n\n## H2\n\n### H3", tmp_path / "c.docx")
    doc = _doc(p)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        pf = doc.styles[name].paragraph_format
        assert pf.space_before is not None and pf.space_before.pt > 0
        assert pf.space_after is not None and pf.space_after.pt > 0


def _style_first_line_chars(doc, style_name):
    """``w:ind/@w:firstLineChars`` on a paragraph style (None if absent)."""
    pPr = doc.styles[style_name].element.find(qn("w:pPr"))
    if pPr is None:
        return None
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        return None
    return ind.get(qn("w:firstLineChars"))


def test_heading_no_first_line_indent(tmp_path):
    # 9.1: Heading styles explicitly reset firstLineChars to 0 — they must not
    # inherit Normal's 2-char body indent (which shifts centered H1 off-center
    # and gives H2/H3 a spurious left indent).
    p = create_docx("# H1\n\n## H2\n\n### H3", tmp_path / "c.docx")
    doc = _doc(p)
    assert _style_first_line_chars(doc, "Normal") == "200"  # body keeps 2-char indent
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        assert _style_first_line_chars(doc, name) == "0", (
            f"{name} leaks Normal's firstLineChars (heading indent bug)"
        )


def test_slot_style_literal_chinese(tmp_path):
    # 9.2: slotted variant renders {{中文slot}} as plain text with NO highlight.
    # Tokens stay find_slots()-fillable (literal preserves {{token}} text).
    p = create_docx("甲方：{{甲方名称}}", tmp_path / "c.docx", slot_style="literal")
    para = _doc(p).paragraphs[0]
    assert "{{甲方名称}}" in para.text
    for r in para.runs:
        rPr = r._r.find(qn("w:rPr"))
        assert rPr is None or rPr.find(qn("w:highlight")) is None
    assert find_slots(p) == ["甲方名称"]


def test_date_line_underlines(tmp_path):
    # 9.3: 日期：______年____月____日 — the _年_/_月_ markdown-underline spans
    # must NOT eat the boundary underscores. Result is three underlined blanks
    # (6/4/4 spaces) and zero literal `_` characters.
    p = create_docx("日期：______年____月____日", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    assert "_" not in para.text  # no literal underscores survived
    blanks = [len(r.text) for r in _underlined_runs(para)]
    assert blanks == [6, 4, 4]  # 年/月/日 are plain text, not underlined spans


def _all_paragraphs(doc):
    # body paragraphs + every paragraph inside table cells (signature block).
    paras = list(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                paras.extend(cell.paragraphs)
    return paras


def test_signature_placeholder_dropped(tmp_path):
    # 10: the 签署信息 placeholder sentence must not render; a default 甲方/乙方
    # multi-line block appears in its place (the sale-base bug).
    content = (
        "## 签署信息\n\n"
        "本条款为合同签署页格式元素，已移至合同签署页单独处理，此处无需保留正文内容。"
    )
    doc = _doc(create_docx(content, tmp_path / "c.docx"))
    texts = [p.text for p in _all_paragraphs(doc)]
    assert not any("本条款为合同签署页格式元素" in t for t in texts)
    joined = "\n".join(texts)
    assert "甲方" in joined and "乙方" in joined  # default block rendered


def test_signature_comma_split_multiline(tmp_path):
    # 10: a comma-separated signature line splits into one field per line
    # (甲方：a，乙方：b，签订日期：d -> 3 separate paragraphs).
    content = (
        "## 签署信息\n\n"
        "甲方（签字/盖章）：{{party_a_sign}}，乙方（签字/盖章）：{{party_b_sign}}，"
        "签订日期：{{sign_date}}"
    )
    doc = _doc(create_docx(content, tmp_path / "c.docx", slot_style="literal"))
    texts = [p.text for p in _all_paragraphs(doc)]
    joined = "\n".join(texts)
    assert "甲方" in joined and "乙方" in joined and "签订日期" in joined
    # no single paragraph crams all three fields on one line (split happened)
    for t in texts:
        assert not ("甲方" in t and "乙方" in t and "签订日期" in t)
