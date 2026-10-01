"""Tests for the corpus reader (src/clauses/corpus.py).

Mostly hermetic (parse/clean on local files + ConfigError on unreachable MinIO);
one integration test exercises the real MinIO corpus and skips if unreachable.
"""

from __future__ import annotations

import pytest

from src.clauses import corpus
from src.settings import ConfigError


def test_clean_strips_boilerplate():
    text = "合同编号：\n使 用 说 明\n\n甲方：{{party_a}}\n\n\n\n乙方：{{party_b}}"
    cleaned = corpus._clean(text)
    assert "合同编号" not in cleaned
    assert "使 用 说 明" not in cleaned
    assert "{{party_a}}" in cleaned
    assert "\n\n\n" not in cleaned  # whitespace runs collapsed


def test_parse_docx_local(tmp_path):
    from docx import Document

    p = tmp_path / "x.docx"
    d = Document()
    d.add_paragraph("甲方：张三")
    d.add_paragraph("乙方：李四")
    d.save(str(p))
    text = corpus._parse_docx(str(p))
    assert "张三" in text and "李四" in text


def test_list_corpus_documents_raises_when_unreachable(monkeypatch):
    monkeypatch.setattr(corpus, "MINIO_ENDPOINT", "")
    with pytest.raises(ConfigError):
        corpus.list_corpus_documents()


def test_corpus_integration_real_minio():
    """Integration: list + read one doc from the real MinIO corpus (skip if unreachable)."""
    try:
        docs = corpus.list_corpus_documents()
    except ConfigError as exc:
        pytest.skip(f"MinIO unreachable: {exc}")
    assert len(docs) > 0
    # read the first docx (pdf parsing also works, but docx is cleaner)
    docx_doc = next((d for d in docs if d.ext == ".docx"), docs[0])
    r = corpus.read_document(docx_doc.object_name)
    assert r["text"].strip() != ""
