"""Tests for the DOCX generator (``src.generator.generator``)."""

from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from src.generator import add_slot, create_docx, find_slots


def _doc(path: Path) -> Document:
    return Document(str(path))


def test_create_docx_renders_headings_lists_paragraphs(tmp_path):
    p = create_docx(
        "# 服务协议\n\n甲方与乙方达成如下协议。\n\n## 第一条 服务内容\n\n"
        "- 甲方提供咨询\n- 乙方支付费用",
        tmp_path / "out" / "contract.docx",
    )
    assert p == tmp_path / "out" / "contract.docx"
    assert p.exists()
    paras = _doc(p).paragraphs
    styles = [para.style.name for para in paras]
    texts = [para.text for para in paras]
    assert styles == ["Heading 1", "Normal", "Heading 2", "List Bullet", "List Bullet"]
    assert texts[0] == "服务协议"
    assert texts[1] == "甲方与乙方达成如下协议。"
    assert texts[3] == "甲方提供咨询" and texts[4] == "乙方支付费用"


def test_create_docx_creates_missing_parent_dirs(tmp_path):
    p = create_docx("# Title", tmp_path / "a" / "b" / "c.docx")
    assert p.exists()


def test_create_docx_title_becomes_first_heading(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx", title="合同标题")
    paras = _doc(p).paragraphs
    assert paras[0].style.name == "Heading 1"
    assert paras[0].text == "合同标题"
    assert paras[1].text == "正文"


def test_create_docx_fills_slots_from_mapping(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx", slots={"party_a": "北京甲公司"})
    assert find_slots(p) == []
    assert _doc(p).paragraphs[0].text == "甲方：北京甲公司"


def test_create_docx_leaves_highlighted_placeholder_when_absent(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx")
    para = _doc(p).paragraphs[0]
    assert "{{party_a}}" in para.text
    highlighted = [
        r.text for r in para.runs
        if r._r.find(qn("w:rPr")) is not None
        and r._r.find(qn("w:rPr")).find(qn("w:highlight")) is not None
    ]
    assert "{{party_a}}" in highlighted


def test_create_docx_explicit_none_leaves_placeholder(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx", slots={"party_a": None})
    assert find_slots(p) == ["party_a"]


def test_create_docx_bold_and_italic(tmp_path):
    p = create_docx("**重要**：*双方*遵守。", tmp_path / "c.docx")
    runs = _doc(p).paragraphs[0].runs
    assert runs[0].text == "重要" and runs[0].bold is True
    assert runs[2].text == "双方" and runs[2].italic is True


def test_create_docx_horizontal_rule(tmp_path):
    p = create_docx("上文\n\n---\n\n下文", tmp_path / "c.docx")
    paras = _doc(p).paragraphs
    hr = paras[1]
    assert hr.text == ""
    pPr = hr._p.find(qn("w:pPr"))
    assert pPr is not None and pPr.find(qn("w:pBdr")) is not None


def test_add_slot_labeled_placeholder(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    add_slot(p, "sign_date", label="签署日期")
    assert _doc(p).paragraphs[-1].text == "签署日期: {{sign_date}}"
    assert find_slots(p) == ["sign_date"]


def test_add_slot_unlabeled_placeholder(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    add_slot(p, "note")
    assert _doc(p).paragraphs[-1].text == "{{note}}"


def test_add_slot_prefilled_value(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    add_slot(p, "amount", label="金额", value="人民币壹万元整")
    assert _doc(p).paragraphs[-1].text == "金额: 人民币壹万元整"
    assert "amount" not in find_slots(p)


def test_add_slot_accepts_open_document(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    doc = Document(str(p))
    add_slot(doc, "x", label="L")
    doc.save(str(p))
    assert _doc(p).paragraphs[-1].text == "L: {{x}}"
