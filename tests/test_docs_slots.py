"""Tests for the slot model (``src.generator.slots``)."""

from pathlib import Path

import pytest
from docx import Document
from docx.oxml.ns import qn

from src.generator import create_docx, fill_slots, find_slots, slot_token, SlotValueError


def test_slot_token_builds_placeholder():
    assert slot_token("party_a") == "{{party_a}}"


def test_slot_token_rejects_invalid_name():
    with pytest.raises(ValueError):
        slot_token("bad name")


def test_find_slots_dedups_in_first_appearance_order(tmp_path):
    p = create_docx(
        "甲方：{{party_a}}\n乙方：{{party_b}}\n日期：{{party_a}}",
        tmp_path / "c.docx",
    )
    assert find_slots(p) == ["party_a", "party_b"]


def test_find_slots_accepts_open_document(tmp_path):
    p = create_docx("{{x}}", tmp_path / "c.docx")
    assert find_slots(Document(str(p))) == ["x"]


def test_fill_slots_replaces_token(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx")
    count = fill_slots(p, {"party_a": "北京甲公司"})
    assert count == 1
    assert find_slots(p) == []
    assert Document(str(p)).paragraphs[0].text == "甲方：北京甲公司"


def test_fill_slots_partial_leaves_others(tmp_path):
    p = create_docx("{{party_a}}\n{{party_b}}", tmp_path / "c.docx")
    count = fill_slots(p, {"party_a": "甲"})
    assert count == 1
    assert find_slots(p) == ["party_b"]


def test_fill_slots_strict_raises_on_missing(tmp_path):
    p = create_docx("{{party_a}}\n{{party_b}}", tmp_path / "c.docx")
    with pytest.raises(SlotValueError) as exc:
        fill_slots(p, {"party_a": "甲"}, strict=True)
    assert "party_b" in str(exc.value)


def test_fill_slots_strict_passes_when_all_present(tmp_path):
    p = create_docx("{{a}}", tmp_path / "c.docx")
    assert fill_slots(p, {"a": "v"}, strict=True) == 1
    assert find_slots(p) == []


def test_fill_slots_multiline_inserts_break(tmp_path):
    p = create_docx("{{a}}", tmp_path / "c.docx")
    fill_slots(p, {"a": "line1\nline2"})
    para = Document(str(p)).paragraphs[0]
    assert para.runs[0]._r.findall(qn("w:br")), "expected a <w:br/> line break"
    assert "line1" in para.text and "line2" in para.text


def test_fill_slots_accepts_open_document(tmp_path):
    p = create_docx("{{a}}", tmp_path / "c.docx")
    doc = Document(str(p))
    assert fill_slots(doc, {"a": "v"}) == 1
    doc.save(str(p))  # caller-owned document must be saved explicitly
    assert find_slots(p) == []
