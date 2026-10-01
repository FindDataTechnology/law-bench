"""Tests for Chinese contract DOCX formatting (``add-contract-docx-formatting``)."""

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from src.generator import create_docx, find_slots


def _doc(path):
    return Document(str(path))


def _rfonts_eastasia(style):
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    return rFonts.get(qn("w:eastAsia"))


def test_normal_style_sets_cjk_font(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    normal = _doc(p).styles["Normal"]
    assert _rfonts_eastasia(normal) == "宋体"


def test_page_is_a4(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    section = _doc(p).sections[0]
    # page_width/height round-trip through twips (1 twip = 635 EMU), so 210mm
    # is not an exact EMU value; allow a sub-mm tolerance.
    assert abs(section.page_width.mm - 210) < 0.1
    assert abs(section.page_height.mm - 297) < 0.1


def test_normal_paragraph_has_two_char_indent(tmp_path):
    p = create_docx("正文", tmp_path / "c.docx")
    normal = _doc(p).styles["Normal"]
    pPr = normal.element.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    assert ind is not None
    assert ind.get(qn("w:firstLineChars")) == "200"


def test_heading1_is_black_cjk_centered(tmp_path):
    p = create_docx("# 标题", tmp_path / "c.docx")
    doc = _doc(p)
    style = doc.styles["Heading 1"]
    assert _rfonts_eastasia(style) == "黑体"
    assert style.font.color.rgb is not None and str(style.font.color.rgb) == "000000"
    # centering is set on the style; the paragraph inherits it (its own
    # alignment stays None until explicitly overridden).
    assert style.paragraph_format.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_highlight_off_removes_highlight_keeps_token(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx", slot_style="literal")
    para = _doc(p).paragraphs[0]
    assert "{{party_a}}" in para.text
    for r in para.runs:
        rPr = r._r.find(qn("w:rPr"))
        assert rPr is None or rPr.find(qn("w:highlight")) is None
    assert find_slots(p) == ["party_a"]


def test_highlight_on_default(tmp_path):
    p = create_docx("甲方：{{party_a}}", tmp_path / "c.docx", slot_style="highlight")
    para = _doc(p).paragraphs[0]
    highlighted = [
        r.text for r in para.runs
        if r._r.find(qn("w:rPr")) is not None
        and r._r.find(qn("w:rPr")).find(qn("w:highlight")) is not None
    ]
    assert "{{party_a}}" in highlighted


def test_signature_section_becomes_table(tmp_path):
    p = create_docx(
        "## 签署信息\n\n{{party_a_sign}}\n\n{{party_b_sign}}",
        tmp_path / "c.docx",
    )
    doc = _doc(p)
    assert len(doc.tables) == 1
    assert len(doc.tables[0].columns) == 2
