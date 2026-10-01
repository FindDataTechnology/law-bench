"""DOCX generation from Markdown contract content with slot support.

The renderer is intentionally lightweight: it covers the Markdown structures
that appear in the project's contract templates (``templates/*.md``) -
headings, bold/italic, bullet and numbered lists, horizontal rules, and
paragraphs. Slot tokens (``{{slot_name}}``) in the content are rendered as
highlighted placeholders or pre-filled values.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from .slots import (
    SLOT_PATTERN,
    _open_document,
    set_run_text,
    slot_token,
)

__all__ = ["create_docx", "add_slot"]

# ponytail: 宋体/黑体 read correctly in Word/Windows; the container's LibreOffice
# needs matching CJK glyphs installed (or a Noto fallback) for PDF. See tasks 6.1.
_BODY_FONT = "宋体"
_HEADING_FONT = "黑体"
_BODY_SIZE = Pt(12)  # 小四
_HEADING_SIZES = {"Heading 1": Pt(22), "Heading 2": Pt(16), "Heading 3": Pt(14)}
_HEADING_SPACING = {"Heading 1": (12, 6), "Heading 2": (8, 4), "Heading 3": (6, 3)}
_LINE_SPACING = 1.5
_FIRST_LINE_CHARS = 200  # 2 chars, in hundredths of a char

_HR_RE = re.compile(r"^(?:-{3,}|\*{3,}|_{3,})$")
_HEADING_RE = re.compile(r"^(#{1,3})\s+(.*)$")
_BULLET_RE = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED_RE = re.compile(r"^\d+\.\s+(.*)$")
# Inline tokens: a slot, a **bold** span, or an *italic* span. The slot name
# pattern mirrors ``SLOT_PATTERN`` (Unicode ``\w``) so CJK slots like
# ``{{甲方名称}}`` are recognized as slots (highlighted/fillable), not just the
# Latin ones (``{{party_a}}``).
_UNDERLINE_RE = re.compile(r"_{3,}")  # literal ____ blanks -> real Word underline
_INLINE_RE = re.compile(
    r"\{\{\s*[\w][\w-]*\s*\}\}"  # slot
    r"|\*\*[^*]+\*\*"  # **bold**
    r"|\*[^*]+\*"  # *italic*
    r"|(?<!_)_[^_]+_(?!_)"  # _underline_ — lookaround guards ___ blank boundaries (9.3)
)


def _parse_inline(text: str) -> list[dict]:
    """Split ``text`` into ordered segments: plain, slot, bold, italic, or underline."""
    segments: list[dict] = []
    pos = 0
    for m in _INLINE_RE.finditer(text):
        if m.start() > pos:
            segments.append({"text": text[pos:m.start()], "bold": False, "italic": False, "underline": False, "slot": None})
        tok = m.group(0)
        if tok.startswith("{{"):
            segments.append(
                {"text": tok, "bold": False, "italic": False, "underline": False, "slot": SLOT_PATTERN.match(tok).group(1)}
            )
        elif tok.startswith("**"):
            segments.append({"text": tok[2:-2], "bold": True, "italic": False, "underline": False, "slot": None})
        elif tok.startswith("_"):  # _underline_
            segments.append({"text": tok[1:-1], "bold": False, "italic": False, "underline": True, "slot": None})
        else:  # *italic*
            segments.append({"text": tok[1:-1], "bold": False, "italic": True, "underline": False, "slot": None})
        pos = m.end()
    if pos < len(text):
        segments.append({"text": text[pos:], "bold": False, "italic": False, "underline": False, "slot": None})
    return segments


def _apply_inline(run, seg: dict) -> None:
    if seg["bold"]:
        run.bold = True
    if seg["italic"]:
        run.italic = True
    if seg.get("underline"):
        run.underline = True


def _blank_width(name: str) -> int:
    # ponytail: fixed multiplier; tune if blanks look too short/long.
    return min(24, max(8, len(name) * 2))


def _split_underlines(text: str) -> list[tuple[str, bool]]:
    """Split plain text into ``(chunk, is_blank)`` pairs; blank runs are ``____``."""
    parts: list[tuple[str, bool]] = []
    pos = 0
    for m in _UNDERLINE_RE.finditer(text):
        if m.start() > pos:
            parts.append((text[pos:m.start()], False))
        parts.append((m.group(0), True))
        pos = m.end()
    if pos < len(text):
        parts.append((text[pos:], False))
    return parts


def _render_inline(paragraph: Paragraph, text: str, slots, slot_style: str = "highlight") -> None:
    """Add runs for ``text``, rendering slots per ``slot_style``."""
    for seg in _parse_inline(text):
        if seg["slot"] is not None:
            name = seg["slot"]
            value = None if slots is None else slots.get(name)
            if value is not None:
                run = paragraph.add_run()
                _apply_inline(run, seg)
                set_run_text(run, str(value))
            elif slot_style == "underline":
                run = paragraph.add_run(" " * _blank_width(name))
                _apply_inline(run, seg)
                run.font.underline = True
            else:
                run = paragraph.add_run(slot_token(name))
                if slot_style == "highlight":
                    run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                _apply_inline(run, seg)
        elif seg["text"]:
            for chunk, is_blank in _split_underlines(seg["text"]):
                run = paragraph.add_run(" " * len(chunk) if is_blank else chunk)
                _apply_inline(run, seg)
                if is_blank:
                    run.font.underline = True


def _add_horizontal_rule(document) -> None:
    paragraph = document.add_paragraph()
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "auto")
    pbdr.append(bottom)
    pPr.append(pbdr)


def _set_eastasia_font(run_or_style, name: str) -> None:
    """Set both Latin and CJK faces on a run or style.

    ``.font.name`` only writes ``w:rFonts/@w:ascii``+``@w:hAnsi`` (Latin); CJK
    glyphs read ``@w:eastAsia``, which must be set separately or Chinese text
    keeps the default face.
    """
    run_or_style.font.name = name  # ascii + hAnsi
    rPr = run_or_style.element.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def _set_first_line_chars(style, chars: int = _FIRST_LINE_CHARS) -> None:
    """2-char first-line indent (scales with font size, unlike fixed twips)."""
    pPr = style.element.get_or_add_pPr()
    ind = pPr.get_or_add_ind()
    ind.set(qn("w:firstLineChars"), str(chars))
    ind.set(qn("w:firstLine"), str(chars * 240 // 100))  # twips fallback


def _set_heading(style, name: str, size, center: bool, spacing: tuple[int, int]) -> None:
    style.font.bold = False  # 黑体 is already a heavy face
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = size
    _set_eastasia_font(style, name)
    # Headings inherit Normal (which carries the 2-char firstLineChars); reset to 0
    # so H1 centers cleanly and H2/H3 are flush-left, not indented like body text.
    _set_first_line_chars(style, 0)
    if center:
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    before, after = spacing
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)


def _apply_zh_document_style(document) -> None:
    """A4 page, 宋体 body with 2-char indent + line spacing, 黑体 headings."""
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Mm(37)
    section.bottom_margin = Mm(35)
    section.left_margin = Mm(28)
    section.right_margin = Mm(26)

    normal = document.styles["Normal"]
    normal.font.size = _BODY_SIZE
    _set_eastasia_font(normal, _BODY_FONT)
    normal.paragraph_format.line_spacing = _LINE_SPACING
    _set_first_line_chars(normal)

    for i, heading in enumerate(("Heading 1", "Heading 2", "Heading 3"), start=1):
        _set_heading(
            document.styles[heading], _HEADING_FONT, _HEADING_SIZES[heading],
            center=(i == 1), spacing=_HEADING_SPACING[heading],
        )


def _set_table_borderless(table) -> None:
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "none")
        borders.append(el)
    table._tbl.tblPr.append(borders)


# ponytail: split single-line comma-separated signatures into multi-line fields
# (甲方：a，乙方：b，日期：d -> 3 lines) so each renders on its own line.
_SIG_SPLIT_RE = re.compile(r"[，；;,]")
# Attribute nouns: a field bearing one attaches to the current party block
# (法定代表人/日期 line under 甲方（盖章）). Avoids enumerating every party noun.
_ATTR_RE = re.compile(
    r"法定代表人|授权代表|委托代理人|代理人|签订日期|签署日期|日期|时间|"
    r"签订地点|签署地点|地点|住所|联系方式|联系电话|电话|地址|身份证号|身份证|邮编|开户行|账号"
)
# Placeholder/meta lines with no real signature -> drop, render default block.
_DROP_SIG_RE = re.compile(r"已移至合同签署页|此处无需保留|无签署占位符")
# Date blanks: "年 月 日" (spaces, no digits) -> underlined blanks.
_DATE_BLANK_RE = re.compile(r"(?<=[:：])\s*年\s*月\s*日")

_DEFAULT_SIG_FIELDS = [
    "甲方（签字/盖章）：________________",
    "日期：____年____月____日",
    "乙方（签字/盖章）：________________",
    "日期：____年____月____日",
]


def _signature_fields(lines: list[str]) -> list[str]:
    """Flatten signature lines into one field per line (多行显示).

    Splits comma-separated signature lines (``甲方：a，乙方：b，日期：d``)
    into separate fields, fills ``年 月 日`` date blanks with underlines, and
    drops the placeholder sentence / ``（本类合同无签署占位符）`` meta notes.
    Falls back to a default 甲方/乙方 block when no real signature remains.
    """
    fields: list[str] = []
    for ln in lines:
        if _DROP_SIG_RE.search(ln):
            continue
        for part in _SIG_SPLIT_RE.split(ln):
            part = part.strip()
            if not part:
                continue
            part = _DATE_BLANK_RE.sub("____年____月____日", part)
            fields.append(part)
    return fields or _DEFAULT_SIG_FIELDS


def _signature_cells(fields: list[str]) -> list[list[str]]:
    """Group fields into party blocks (one cell per party), multi-line."""
    cells: list[list[str]] = []
    cur: list[str] = []
    for ln in fields:
        if cur and not _ATTR_RE.search(ln):
            cells.append(cur)
            cur = [ln]
        else:
            cur.append(ln)
    if cur:
        cells.append(cur)
    return cells


def _fill_signature_cell(cell, lines, slots, slot_style) -> None:
    first = True
    for ln in lines:
        para = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        para.alignment = WD_ALIGN_PARAGRAPH.LEFT
        _render_inline(para, ln, slots, slot_style)


def _render_signature_block(document, lines, slots, slot_style) -> None:
    """Render 签署信息 as a borderless 2-column, left-aligned multi-line table.

    Each party block (甲方 + its 代理人/日期 fields) is one multi-line cell;
    comma-separated single-line signatures are split so each field gets its own
    line. Placeholder-only sections fall back to a default 甲方/乙方 block.
    """
    fields = _signature_fields(lines)
    cells = _signature_cells(fields)
    table = document.add_table(rows=0, cols=2)
    _set_table_borderless(table)
    for k in range(0, len(cells), 2):
        left = cells[k]
        right = cells[k + 1] if k + 1 < len(cells) else []
        row = table.add_row()
        _fill_signature_cell(row.cells[0], left, slots, slot_style)
        _fill_signature_cell(row.cells[1], right, slots, slot_style)


def _collect_signature_lines(lines: list[str], start: int) -> tuple[list[str], int]:
    """Collect stripped, non-empty lines from ``start`` until a heading/HR."""
    sig: list[str] = []
    i = start
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if _HEADING_RE.match(line) or _HR_RE.match(line):
            break
        sig.append(line)
        i += 1
    return sig, i


def _render_markdown(document, content: str, slots, slot_style: str = "highlight") -> None:
    lines = content.splitlines()
    i = 0
    n = len(lines)
    while i < n:
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if _HR_RE.match(line):
            _add_horizontal_rule(document)
            i += 1
        elif m := _HEADING_RE.match(line):
            is_signature = "签署信息" in m.group(2)
            paragraph = document.add_paragraph(style=f"Heading {len(m.group(1))}")
            _render_inline(paragraph, m.group(2), slots, slot_style)
            i += 1
            if is_signature:
                sig_lines, i = _collect_signature_lines(lines, i)
                if sig_lines:
                    _render_signature_block(document, sig_lines, slots, slot_style)
        elif "签字/盖章" in line:
            sig_lines, i = _collect_signature_lines(lines, i)
            _render_signature_block(document, sig_lines, slots, slot_style)
        elif m := _BULLET_RE.match(line):
            paragraph = document.add_paragraph(style="List Bullet")
            _render_inline(paragraph, m.group(1), slots, slot_style)
            i += 1
        elif m := _NUMBERED_RE.match(line):
            paragraph = document.add_paragraph(style="List Number")
            _render_inline(paragraph, m.group(1), slots, slot_style)
            i += 1
        else:
            paragraph = document.add_paragraph()
            _render_inline(paragraph, line, slots, slot_style)
            i += 1


def create_docx(content, output_path, *, slots=None, title=None, slot_style: str = "highlight"):
    """Render Markdown ``content`` to a ``.docx`` at ``output_path``.

    ``slots`` (optional) maps slot names to values: a non-None value fills the
    ``{{slot}}`` token at generation time; a missing or None entry leaves a
    placeholder. ``title`` (optional) is rendered as a Heading 1 before the
    content. ``slot_style`` controls how unfilled ``{{slot}}`` tokens render:
    ``"highlight"`` (default) yellow-highlights the token text; ``"underline"``
    renders an N-space underlined blank with no token text (terminal 成品);
    ``"literal"`` renders the token text with no highlight (legacy behavior).
    Parent directories are created as needed. Returns the resolved
    :class:`~pathlib.Path` to the written file.
    """
    if slot_style not in ("highlight", "underline", "literal"):
        raise ValueError(
            f"unknown slot_style: {slot_style!r}; use highlight, underline, or literal"
        )
    document = Document()
    _apply_zh_document_style(document)
    if title:
        heading = document.add_paragraph(style="Heading 1")
        _render_inline(heading, str(title), slots, slot_style)
    _render_markdown(document, content, slots, slot_style)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(out))
    return out


def add_slot(doc, slot_name, *, label=None, value=None):
    """Append a slot paragraph to ``doc`` (a path or open document).

    With ``label`` the paragraph reads ``label: {{slot_name}}`` (or just
    ``{{slot_name}}``). The placeholder is highlighted unless ``value`` is
    given, in which case the slot is filled immediately. When ``doc`` is a path
    the file is saved in place.
    """
    document = _open_document(doc)
    paragraph = document.add_paragraph()
    if label:
        paragraph.add_run(f"{label}: ")
    if value is not None:
        run = paragraph.add_run()
        set_run_text(run, str(value))
    else:
        run = paragraph.add_run(slot_token(slot_name))
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    if isinstance(doc, (str, Path)):
        document.save(str(doc))
    return None
